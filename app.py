from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
import os
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill
from werkzeug.security import generate_password_hash, check_password_hash
import psycopg2
from psycopg2.extras import DictCursor

app = Flask(__name__)
app.secret_key = 'clave_secreta_super_segura'

# ⚠️ TU BASE DE DATOS CONECTADA AQUÍ
DATABASE_URL = "postgresql://db_inventario_n717_user:LQ7bQFWc4r64VqUPGNif7AmZSD9b5eu4@dpg-dbcdrafavr4c73fmhhjg-a/db_inventario_n717"

def get_db():
    conn = psycopg2.connect(DATABASE_URL, sslmode='require')
    return conn

def init_db():
    conn = get_db()
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS productos (
        id SERIAL PRIMARY KEY,
        nombre TEXT NOT NULL,
        categoria TEXT,
        cantidad INTEGER DEFAULT 0,
        stock_minimo INTEGER DEFAULT 5,
        precio REAL DEFAULT 0,
        foto TEXT,
        fecha_actualizacion TEXT
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS usuarios (
        id SERIAL PRIMARY KEY,
        username TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL,
        rol TEXT NOT NULL
    )''')
    try:
        c.execute("INSERT INTO usuarios (username, password, rol) VALUES (%s, %s, %s)",
                  ('admin', generate_password_hash('admin123'), 'admin'))
        c.execute("INSERT INTO usuarios (username, password, rol) VALUES (%s, %s, %s)",
                  ('trabajador', generate_password_hash('trabajador123'), 'trabajador'))
    except Exception:
        pass
    conn.commit()
    c.close()
    conn.close()

# ================= RUTAS =================

@app.route('/instalar_tablas')
def instalar_tablas():
    try:
        init_db()
        return "<h1>¡Tablas creadas con éxito!</h1><p>Ya puedes ir a la página principal y loguearte.</p>"
    except Exception as e:
        return f"<h1>Error:</h1><p>{str(e)}</p>"

@app.route('/')
def index():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    return redirect(url_for('menu'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        conn = get_db()
        c = conn.cursor(cursor_factory=DictCursor)
        c.execute("SELECT * FROM usuarios WHERE username = %s", (username,))
        user = c.fetchone()
        c.close()
        conn.close()
        
        if user and check_password_hash(user['password'], password):
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['rol'] = user['rol']
            return redirect(url_for('menu'))
        else:
            flash('Usuario o contraseña incorrectos', 'danger')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/menu')
def menu():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    return render_template('menu.html')

@app.route('/stock')
def stock():
    if session.get('rol') != 'admin':
        return redirect(url_for('menu'))
    conn = get_db()
    c = conn.cursor(cursor_factory=DictCursor)
    c.execute("SELECT * FROM productos ORDER BY nombre")
    productos = c.fetchall()
    c.close()
    conn.close()
    return render_template('stock.html', productos=productos)

@app.route('/alertas')
def alertas():
    if session.get('rol') != 'admin':
        return redirect(url_for('menu'))
    conn = get_db()
    c = conn.cursor(cursor_factory=DictCursor)
    c.execute("SELECT * FROM productos WHERE cantidad <= stock_minimo ORDER BY nombre")
    productos = c.fetchall()
    c.close()
    conn.close()
    return render_template('alertas.html', productos=productos)

@app.route('/form_ingreso')
def form_ingreso():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    conn = get_db()
    c = conn.cursor(cursor_factory=DictCursor)
    c.execute("SELECT * FROM productos ORDER BY nombre")
    productos = c.fetchall()
    c.execute("SELECT DISTINCT categoria FROM productos WHERE categoria IS NOT NULL AND categoria != '' ORDER BY categoria")
    categorias = [row['categoria'] for row in c.fetchall()]
    c.close()
    conn.close()
    fecha_hoy = datetime.now().strftime('%Y-%m-%d')
    return render_template('form_ingreso.html', productos=productos, categorias=categorias, fecha_hoy=fecha_hoy)

@app.route('/form_venta')
def form_venta():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    conn = get_db()
    c = conn.cursor(cursor_factory=DictCursor)
    c.execute("SELECT * FROM productos ORDER BY nombre")
    productos = c.fetchall()
    c.execute("SELECT DISTINCT categoria FROM productos WHERE categoria IS NOT NULL AND categoria != '' ORDER BY categoria")
    categorias = [row['categoria'] for row in c.fetchall()]
    c.close()
    conn.close()
    fecha_hoy = datetime.now().strftime('%Y-%m-%d')
    return render_template('form_venta.html', productos=productos, categorias=categorias, fecha_hoy=fecha_hoy)

@app.route('/accion', methods=['POST'])
def accion():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    try:
        producto_id = request.form.get('producto_id')
        cantidad = request.form.get('cantidad')
        tipo = request.form.get('tipo')
        fecha_formulario = request.form.get('fecha_accion')
        metodo_pago = request.form.get('metodo_pago', '')
        precio_venta = request.form.get('precio_venta')
        
        if not producto_id or not cantidad:
            flash('Error: Faltan datos.', 'danger')
            return redirect(url_for('menu'))

        cantidad = int(cantidad)
        conn = get_db()
        c = conn.cursor(cursor_factory=DictCursor)
        c.execute("SELECT * FROM productos WHERE id = %s", (producto_id,))
        p = c.fetchone()
        
        if p:
            if tipo == 'venta' and p['cantidad'] < cantidad:
                flash(f'Error: Solo hay {p["cantidad"]} en stock.', 'danger')
            else:
                nueva_cantidad = p['cantidad'] + cantidad if tipo == 'ingreso' else p['cantidad'] - cantidad
                fecha = datetime.strptime(fecha_formulario, '%Y-%m-%d').strftime("%Y-%m-%d") if fecha_formulario else datetime.now().strftime("%Y-%m-%d")
                hora_actual = datetime.now().strftime("%H:%M")
                
                if tipo == 'venta':
                    precio_final = float(precio_venta) if precio_venta else p['precio']
                    fecha_completa = f"{fecha} {hora_actual} | Cant: {cantidad} | Precio: ${precio_final} | Pago: {metodo_pago}"
                    c.execute("UPDATE productos SET precio = %s WHERE id = %s", (precio_final, producto_id))
                else:
                    fecha_completa = f"{fecha} {hora_actual} | Ingreso"
                
                c.execute("UPDATE productos SET cantidad = %s, fecha_actualizacion = %s WHERE id = %s",
                             (nueva_cantidad, fecha_completa, producto_id))
                conn.commit()
                
                if nueva_cantidad <= p['stock_minimo']:
                    flash(f'¡{tipo.capitalize()} registrado! ⚠️ Quedan pocas unidades de {p["nombre"]}.', 'warning')
                else:
                    flash(f'¡{tipo.capitalize()} registrado con éxito!', 'success')
        c.close()
        conn.close()
    except Exception as e:
        flash(f'Error: {str(e)}', 'danger')
    return redirect(url_for('menu'))

@app.route('/crear_producto', methods=['POST'])
def crear_producto():
    if session.get('rol') != 'admin':
        return redirect(url_for('login'))
    nombre = request.form['nombre']
    categoria = request.form['categoria']
    cantidad = int(request.form['cantidad'])
    stock_minimo = int(request.form['stock_minimo'])
    precio = float(request.form['precio'])
    fecha = datetime.now().strftime("%Y-%m-%d %H:%M")
    
    conn = get_db()
    c = conn.cursor()
    c.execute("INSERT INTO productos (nombre, categoria, cantidad, stock_minimo, precio, fecha_actualizacion) VALUES (%s, %s, %s, %s, %s, %s)",
                 (nombre, categoria, cantidad, stock_minimo, precio, fecha))
    conn.commit()
    c.close()
    conn.close()
    flash('Producto creado con éxito', 'success')
    return redirect(url_for('stock'))

@app.route('/editar_producto', methods=['POST'])
def editar_producto():
    if session.get('rol') != 'admin':
        return redirect(url_for('login'))
    id_prod = request.form['id_producto']
    nombre = request.form['nombre']
    categoria = request.form['categoria']
    stock_minimo = int(request.form['stock_minimo'])
    precio = float(request.form['precio'])
    
    conn = get_db()
    c = conn.cursor()
    c.execute("UPDATE productos SET nombre=%s, categoria=%s, stock_minimo=%s, precio=%s WHERE id=%s",
                 (nombre, categoria, stock_minimo, precio, id_prod))
    conn.commit()
    c.close()
    conn.close()
    flash('Producto actualizado con éxito', 'success')
    return redirect(url_for('stock'))

@app.route('/eliminar_producto/<int:id>')
def eliminar_producto(id):
    if session.get('rol') != 'admin':
        return redirect(url_for('login'))
    conn = get_db()
    c = conn.cursor()
    c.execute("DELETE FROM productos WHERE id=%s", (id,))
    conn.commit()
    c.close()
    conn.close()
    flash('Producto eliminado', 'success')
    return redirect(url_for('stock'))

@app.route('/usuarios')
def usuarios():
    if session.get('rol') != 'admin':
        return redirect(url_for('menu'))
    conn = get_db()
    c = conn.cursor(cursor_factory=DictCursor)
    c.execute("SELECT * FROM usuarios ORDER BY id")
    lista_usuarios = c.fetchall()
    c.close()
    conn.close()
    return render_template('usuarios.html', usuarios=lista_usuarios)

@app.route('/crear_usuario', methods=['POST'])
def crear_usuario():
    if session.get('rol') != 'admin':
        return redirect(url_for('menu'))
    username = request.form['username'].strip().lower()
    password = request.form['password']
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM usuarios WHERE username = %s", (username,))
    if c.fetchone():
        flash(f'El usuario "{username}" ya existe.', 'danger')
    else:
        c.execute("INSERT INTO usuarios (username, password, rol) VALUES (%s, %s, %s)",
                     (username, generate_password_hash(password), 'trabajador'))
        conn.commit()
        flash(f'¡Usuario "{username}" creado con éxito!', 'success')
    c.close()
    conn.close()
    return redirect(url_for('usuarios'))

@app.route('/eliminar_usuario/<int:id>')
def eliminar_usuario(id):
    if session.get('rol') != 'admin':
        return redirect(url_for('menu'))
    conn = get_db()
    c = conn.cursor(cursor_factory=DictCursor)
    c.execute("SELECT * FROM usuarios WHERE id = %s", (id,))
    user = c.fetchone()
    if user and user['rol'] != 'admin':
        c.execute("DELETE FROM usuarios WHERE id = %s", (id,))
        conn.commit()
        flash('Usuario eliminado correctamente', 'success')
    else:
        flash('No puedes eliminar a un administrador', 'danger')
    c.close()
    conn.close()
    return redirect(url_for('usuarios'))

@app.route('/exportar_excel')
def exportar_excel():
    if session.get('rol') != 'admin':
        return redirect(url_for('login'))
    conn = get_db()
    c = conn.cursor(cursor_factory=DictCursor)
    c.execute("SELECT * FROM productos ORDER BY nombre")
    productos = c.fetchall()
    c.close()
    conn.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["ID", "Nombre", "Categoría", "Cantidad", "Stock Mínimo", "Precio", "Fecha / Detalle"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="000000")
        
    for row_idx, p in enumerate(productos, 2):
        ws.cell(row=row_idx, column=1, value=p['id'])
        ws.cell(row=row_idx, column=2, value=p['nombre'])
        ws.cell(row=row_idx, column=3, value=p['categoria'])
        ws.cell(row=row_idx, column=4, value=p['cantidad'])
        ws.cell(row=row_idx, column=5, value=p['stock_minimo'])
        ws.cell(row=row_idx, column=6, value=p['precio'])
        ws.cell(row=row_idx, column=7, value=p['fecha_actualizacion'])
        if p['cantidad'] <= p['stock_minimo']:
            for c_idx in range(1, 8):
                ws.cell(row=row_idx, column=c_idx).fill = PatternFill("solid", fgColor="FFD700")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, download_name=f"inventario_{datetime.now().strftime('%Y%m%d')}.xlsx", as_attachment=True)

@app.route('/historial_ventas')
def historial_ventas():
    if session.get('rol') != 'admin':
        return redirect(url_for('menu'))
    conn = get_db()
    c = conn.cursor(cursor_factory=DictCursor)
    c.execute("SELECT * FROM productos WHERE fecha_actualizacion LIKE '%%Pago:%%' ORDER BY fecha_actualizacion DESC")
    ventas = c.fetchall()
    c.close()
    conn.close()
    return render_template('historial_ventas.html', ventas=ventas)

@app.route('/descargar_ventas_fechas', methods=['POST'])
def descargar_ventas_fechas():
    if session.get('rol') != 'admin':
        return redirect(url_for('menu'))
    fecha_inicio = request.form['fecha_inicio']
    fecha_fin = request.form['fecha_fin']
    conn = get_db()
    c = conn.cursor(cursor_factory=DictCursor)
    c.execute("SELECT * FROM productos WHERE fecha_actualizacion LIKE '%%Pago:%%' AND fecha_actualizacion >= %s AND fecha_actualizacion <= %s ORDER BY fecha_actualizacion DESC", 
                          (fecha_inicio, fecha_fin + " 23:59"))
    ventas = c.fetchall()
    c.close()
    conn.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial de Ventas"
    headers = ["Producto", "Categoría", "Detalle de la Venta (Cantidad, Precio y Pago)"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="000000")
        
    for row_idx, p in enumerate(ventas, 2):
        ws.cell(row=row_idx, column=1, value=p['nombre'])
        ws.cell(row=row_idx, column=2, value=p['categoria'])
        ws.cell(row=row_idx, column=3, value=p['fecha_actualizacion'])
        
    for i, w in enumerate([30, 20, 60], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, download_name=f"ventas_{fecha_inicio}_a_{fecha_fin}.xlsx", as_attachment=True)

# ================= INICIALIZACIÓN =================
# Forzamos la creación de las tablas al iniciar
with app.app_context():
    init_db()

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
    
